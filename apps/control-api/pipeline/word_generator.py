"""
pipeline/word_generator.py - Word (.docx) Incident Report Generator

Produces a structured Word document from the scored report dict returned
by run_bridge(). Mirrors the sections shown in the frontend ReportView.
"""
from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# -- Colour palette ------------------------------------------------------------
_BLUE   = RGBColor(0x1a, 0x3a, 0x5c)
_ACCENT = RGBColor(0x4a, 0x86, 0xe8)
_GREEN  = RGBColor(0x43, 0xd3, 0x9f)
_RED    = RGBColor(0xef, 0x53, 0x50)
_AMBER  = RGBColor(0xf2, 0xb9, 0x5b)
_GREY   = RGBColor(0x88, 0x88, 0x88)
_WHITE  = RGBColor(0xff, 0xff, 0xff)
_DARK   = RGBColor(0x1e, 0x22, 0x2b)


def _verdict_color(verdict: Optional[str]) -> RGBColor:
    return {
        "excellent":   _GREEN,
        "acceptable":  _ACCENT,
        "needs_review": _AMBER,
        "failed":      _RED,
    }.get(verdict or "", _GREY)


def _rule_color(status: str) -> RGBColor:
    return {"triggered": _RED, "passed": _GREEN}.get(status, _GREY)


def _fmt_ts(ts: Optional[str]) -> str:
    if not ts:
        return "—"
    try:
        dt = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
        return dt.strftime("%b %d, %Y  %H:%M:%S")
    except Exception:
        return str(ts)


def _fmt_sec(val) -> str:
    if val is None:
        return "—"
    try:
        total = int(float(val))
        m, s = divmod(total, 60)
        h, m = divmod(m, 60)
        if h:
            return f"{h}h {m}m {s}s"
        if m:
            return f"{m}m {s}s"
        return f"{s}s"
    except Exception:
        return str(val)


# -- Document helpers ----------------------------------------------------------

def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    run = p.runs[0] if p.runs else p.add_run(text)
    run.font.color.rgb = _BLUE
    run.font.bold = True


def _kv_row(table, label: str, value: str, value_color: Optional[RGBColor] = None):
    row = table.add_row()
    row.cells[0].text = label
    row.cells[0].paragraphs[0].runs[0].font.bold = True
    row.cells[0].paragraphs[0].runs[0].font.color.rgb = _GREY
    row.cells[1].text = value
    if value_color:
        row.cells[1].paragraphs[0].runs[0].font.color.rgb = value_color


def _metric_table(doc: Document, rows: list[tuple[str, str, Optional[RGBColor]]]):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"
    for cell in hdr:
        _set_cell_bg(cell, "1a3a5c")
        cell.paragraphs[0].runs[0].font.color.rgb = _WHITE
        cell.paragraphs[0].runs[0].font.bold = True
    for label, value, color in rows:
        _kv_row(table, label, value, color)
    doc.add_paragraph()


# -- Section builders ----------------------------------------------------------

def _section_summary(doc: Document, report: dict, member_id: Optional[str]):
    _heading(doc, "Summary", level=1)
    scope = f"Analyst: {member_id}" if member_id else "Team Report"
    rows = [
        ("Incident ID",  report.get("incident_id", "—"), None),
        ("Scenario ID",  report.get("scenario_id", "—"), None),
        ("Scope",        scope, None),
        ("Status",       report.get("status", "—"), _ACCENT),
        ("Outcome",      report.get("outcome", "—"),
         _GREEN if report.get("outcome") == "SUCCESS" else _RED),
    ]
    _metric_table(doc, rows)


def _section_score(doc: Document, report: dict, member_id: Optional[str]):
    _heading(doc, "Score & Verdict", level=1)
    score = report.get("final_score")
    verdict = report.get("verdict")
    label = "Individual Score" if member_id else "Final Score"
    rows = [
        (label,              f"{score} / 100" if score is not None else "—",
         _verdict_color(verdict)),
        ("Verdict",          (verdict or "—").upper(), _verdict_color(verdict)),
        ("Penalty Total",    f"{report.get('penalty_total', 0)} pts",
         _RED if (report.get("penalty_total") or 0) < 0 else None),
        ("Difficulty Bonus", f"+{round(report.get('response_difficulty_bonus', 0), 1)} pts", _GREEN),
        ("TTC Factor",       str(report.get("ttc_factor", "—")), None),
    ]
    if not member_id and report.get("team_score_basis") == "average_of_members":
        ms = report.get("member_scores", {})
        rows.append(("Score Basis",
                      f"Average of {len(ms)} analyst score(s)", _GREY))
    _metric_table(doc, rows)


def _section_timing(doc: Document, report: dict):
    _heading(doc, "Metrics", level=1)
    rows = [
        ("Time to Detect (TTD)",      _fmt_sec(report.get("ttd")),         None),
        ("Time to Contain (TTC)",     _fmt_sec(report.get("ttc")),         None),
        ("Actual TTC (measured)",     _fmt_sec(report.get("ttc_actual_sec")), None),
    ]
    _metric_table(doc, rows)

    _heading(doc, "Timeline", level=1)
    rt = report.get("room_timing", {})
    tl_rows = [
        ("Attack Started",   _fmt_ts(rt.get("attack_started_at")),        None),
        ("Incident Start",   _fmt_ts(report.get("start_time")),           None),
        ("Blue Team Entered",_fmt_ts(rt.get("blue_started_at")),          None),
        ("Detection",        _fmt_ts(report.get("detection_time")),       None),
        ("Scoring Started",  _fmt_ts(report.get("scoring_started_at")),   None),
        ("Containment",      _fmt_ts(report.get("containment_time")),     None),
        ("Incident End",     _fmt_ts(report.get("end_time")),             None),
        ("Exercise Mode",    rt.get("exercise_mode", "—"),                None),
    ]
    _metric_table(doc, tl_rows)


def _section_rules(doc: Document, report: dict):
    _heading(doc, "Score Breakdown (9 Rules)", level=1)
    rules = report.get("scoring_rules", [])
    if not rules:
        doc.add_paragraph("No rule data available.")
        return

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Rule", "Description", "Status", "Penalty"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        _set_cell_bg(hdr_cells[i], "1a3a5c")
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = _WHITE
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True

    for rule in rules:
        row = table.add_row()
        row.cells[0].text = rule.get("rule_id", "")
        row.cells[1].text = rule.get("description", "")
        status = rule.get("status", "")
        row.cells[2].text = {"triggered": "FAIL", "passed": "PASS",
                              "not_applicable": "N/A"}.get(status, status)
        row.cells[2].paragraphs[0].runs[0].font.color.rgb = _rule_color(status)
        pts = rule.get("penalty_points", 0)
        row.cells[3].text = f"{pts} pts" if pts else "—"
        if pts < 0:
            row.cells[3].paragraphs[0].runs[0].font.color.rgb = _RED

        evidence = rule.get("evidence", [])
        if evidence:
            ev_text = "\n".join(f"• {e}" for e in evidence)
            # append evidence as a merged note below
            ev_row = table.add_row()
            merged = ev_row.cells[0].merge(ev_row.cells[3])
            merged.text = ev_text
            merged.paragraphs[0].runs[0].font.size = Pt(8)
            merged.paragraphs[0].runs[0].font.color.rgb = _GREY
            _set_cell_bg(merged, "f5f5f5")

    doc.add_paragraph()


def _section_mitre(doc: Document, report: dict):
    mitre = report.get("mitre_attack", {})
    if not mitre:
        return
    _heading(doc, "MITRE ATT&CK", level=1)
    primary = mitre.get("primary", [])
    related = mitre.get("related", [])

    if primary or related:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for i, h in enumerate(["Technique ID", "Name", "Tactic"]):
            table.rows[0].cells[i].text = h
            _set_cell_bg(table.rows[0].cells[i], "1a3a5c")
            table.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = _WHITE
            table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True

        for t in primary:
            row = table.add_row()
            row.cells[0].text = t.get("technique_id", "")
            row.cells[1].text = t.get("technique_name", "")
            row.cells[2].text = t.get("tactic", "")
            if t.get("source"):
                src_row = table.add_row()
                merged = src_row.cells[0].merge(src_row.cells[2])
                merged.text = t["source"]
                merged.paragraphs[0].runs[0].font.size = Pt(8)
                merged.paragraphs[0].runs[0].font.color.rgb = _ACCENT
                _set_cell_bg(merged, "f0f4ff")

        for t in related:
            row = table.add_row()
            row.cells[0].text = t.get("technique_id", "")
            row.cells[0].paragraphs[0].runs[0].font.color.rgb = _GREY
            row.cells[1].text = f"{t.get('technique_name', '')} (related)"
            row.cells[1].paragraphs[0].runs[0].font.color.rgb = _GREY
            row.cells[2].text = t.get("tactic", "")
            row.cells[2].paragraphs[0].runs[0].font.color.rgb = _GREY
            if t.get("source"):
                src_row = table.add_row()
                merged = src_row.cells[0].merge(src_row.cells[2])
                merged.text = t["source"]
                merged.paragraphs[0].runs[0].font.size = Pt(8)
                merged.paragraphs[0].runs[0].font.color.rgb = _ACCENT
                _set_cell_bg(merged, "f0f4ff")

        doc.add_paragraph()

    if mitre.get("mapping_note"):
        p = doc.add_paragraph()
        p.add_run("Mapping Note: ").bold = True
        p.add_run(mitre["mapping_note"])
        doc.add_paragraph()


def _section_framework(doc: Document, report: dict):
    fw = report.get("response_framework", {})
    if not fw.get("name"):
        return
    _heading(doc, "Response Framework", level=1)
    rows = [
        ("Framework", fw.get("name", "—"), None),
    ]
    if fw.get("note"):
        rows.append(("Note", fw["note"], None))
    if fw.get("source"):
        rows.append(("Reference", fw["source"], _ACCENT))
    _metric_table(doc, rows)


def _section_contributions(doc: Document, report: dict):
    contribs = report.get("analyst_contributions", [])
    if not contribs:
        return
    _heading(doc, "Analyst Contributions", level=1)
    for c in contribs:
        _heading(doc, c.get("analyst_id", "Unknown"), level=2)
        rows = [
            ("First Action", _fmt_ts(c.get("first_action_at")), None),
            ("Last Action",  _fmt_ts(c.get("last_action_at")),  None),
            ("First Action Offset",
             f"+{round(c['first_action_offset_sec'])}s" if c.get("first_action_offset_sec") is not None else "—",
             None),
            ("Total Actions", str(c.get("event_count", 0)), None),
        ]
        et = c.get("event_types", {})
        if et:
            rows.append(("Action Types",
                          ", ".join(f"{k}: {v}" for k, v in et.items()), None))
        _metric_table(doc, rows)


def _section_member_scores(doc: Document, report: dict):
    member_scores = report.get("member_scores", {})
    if not member_scores:
        return
    _heading(doc, "Per-Analyst Scores", level=1)
    for analyst_id, m in member_scores.items():
        _heading(doc, analyst_id, level=2)
        verdict = m.get("verdict")
        rows = [
            ("Score",   f"{m.get('final_score', '—')} / 100", _verdict_color(verdict)),
            ("Verdict", (verdict or "—").upper(),              _verdict_color(verdict)),
            ("Penalty", f"{m.get('penalty_total', 0)} pts",
             _RED if (m.get("penalty_total") or 0) < 0 else None),
        ]
        _metric_table(doc, rows)

        rules = m.get("scoring_rules", [])
        if rules:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Table Grid"
            for i, h in enumerate(["Rule", "Status", "Penalty"]):
                table.rows[0].cells[i].text = h
                _set_cell_bg(table.rows[0].cells[i], "1a3a5c")
                table.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = _WHITE
                table.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
            for rule in rules:
                row = table.add_row()
                row.cells[0].text = rule.get("rule_id", "")
                status = rule.get("status", "")
                row.cells[1].text = {"triggered": "FAIL", "passed": "PASS",
                                      "not_applicable": "N/A"}.get(status, status)
                row.cells[1].paragraphs[0].runs[0].font.color.rgb = _rule_color(status)
                pts = rule.get("penalty_points", 0)
                row.cells[2].text = f"{pts} pts" if pts else "—"
            doc.add_paragraph()


# -- Public entry point --------------------------------------------------------

def generate_word(report: dict, member_id: Optional[str] = None) -> bytes:
    """
    Generate a .docx report from the scored report dict.
    Returns raw bytes ready to serve as application/vnd.openxmlformats-officedocument.wordprocessingml.document.

    member_id — when set, scopes title and score label to one analyst.
    """
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin   = Inches(1.2)
        section.right_margin  = Inches(1.2)

    # Title
    incident_id = report.get("incident_id", "Unknown")
    title_text = (
        f"Incident Report — {incident_id}\nAnalyst: {member_id}"
        if member_id else
        f"Incident Report — {incident_id}"
    )
    title = doc.add_heading(title_text, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = _BLUE
        run.font.size = Pt(18)

    generated_at = datetime.utcnow().strftime("%B %d, %Y at %H:%M UTC")
    sub = doc.add_paragraph(f"Generated: {generated_at}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = _GREY
    sub.runs[0].font.size = Pt(9)
    doc.add_paragraph()

    _section_summary(doc, report, member_id)
    _section_score(doc, report, member_id)
    _section_timing(doc, report)
    _section_rules(doc, report)
    _section_mitre(doc, report)
    _section_framework(doc, report)

    if not member_id:
        _section_contributions(doc, report)
        _section_member_scores(doc, report)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
